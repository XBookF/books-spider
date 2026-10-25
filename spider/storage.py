"""存储层：JSON / CSV / Excel / MySQL（附 SQL 脚本导出，MySQL 不可用时也能演示）"""
import csv
import json
import logging

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config

logger = logging.getLogger(__name__)

FIELDS = ["upc", "title", "category", "price", "price_excl_tax", "tax",
          "rating", "stock", "reviews", "description", "image_url", "detail_url"]

CREATE_TABLE_SQL = """
CREATE TABLE IF NOT EXISTS books (
    id INT AUTO_INCREMENT PRIMARY KEY,
    upc VARCHAR(32) NOT NULL UNIQUE,
    title VARCHAR(255) NOT NULL,
    category VARCHAR(64),
    price DECIMAL(10, 2),
    price_excl_tax DECIMAL(10, 2),
    tax DECIMAL(10, 2),
    rating TINYINT,
    stock INT,
    reviews INT DEFAULT 0,
    description TEXT,
    image_url VARCHAR(512),
    detail_url VARCHAR(512),
    crawl_time TIMESTAMP DEFAULT CURRENT_TIMESTAMP
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
"""


def save_json(records, path):
    path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    logger.info("JSON 已保存：%s（%s 条）", path, len(records))


def save_csv(records, path):
    # utf-8-sig：带 BOM，Excel 双击打开不乱码，这是实际项目里最常见的坑
    with open(path, "w", newline="", encoding="utf-8-sig") as fp:
        writer = csv.DictWriter(fp, fieldnames=FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)
    logger.info("CSV 已保存：%s（%s 条）", path, len(records))


def save_excel(records, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "books"

    header_fill = PatternFill("solid", fgColor="4472C4")
    for col, name in enumerate(FIELDS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, record in enumerate(records, start=2):
        for col_idx, field in enumerate(FIELDS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=record.get(field))

    widths = {"title": 45, "description": 60, "detail_url": 55, "image_url": 55}
    for col_idx, field in enumerate(FIELDS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(field, 14)
    ws.freeze_panes = "A2"                     # 冻结表头
    ws.auto_filter.ref = ws.dimensions         # 开启筛选
    wb.save(path)
    logger.info("Excel 已保存：%s（%s 条）", path, len(records))


def build_insert_sql(records):
    """生成批量插入脚本：ON DUPLICATE KEY UPDATE 依赖 upc 唯一索引实现幂等"""
    lines = ["USE books_db;", CREATE_TABLE_SQL.strip(), ""]
    for r in records:
        values = [
            r["upc"], r["title"].replace("'", "''"), r["category"].replace("'", "''"),
            r["price"], r["price_excl_tax"], r["tax"], r["rating"], r["stock"],
            r["reviews"], r["description"].replace("'", "''"),
            r["image_url"], r["detail_url"],
        ]
        sql = (
            "INSERT INTO books (upc, title, category, price, price_excl_tax, tax, rating, "
            "stock, reviews, description, image_url, detail_url) VALUES ("
            + ", ".join(f"'{v}'" for v in values)
            + ") ON DUPLICATE KEY UPDATE price=VALUES(price), stock=VALUES(stock), "
              "rating=VALUES(rating), description=VALUES(description);"
        )
        lines.append(sql)
    return "\n".join(lines)


def save_sql_file(records, path):
    path.write_text(build_insert_sql(records), encoding="utf-8")
    logger.info("SQL 脚本已保存：%s（可直接在 MySQL 中 source 执行）", path)


def save_mysql(records):
    """写入 MySQL；连不上时返回 False，由调用方降级为导出 SQL 文件"""
    try:
        import pymysql
        from pymysql.err import MySQLError
    except ImportError:
        logger.warning("未安装 pymysql，跳过入库：pip install pymysql")
        return False

    try:
        # 先不指定 database：库还没创建时直接连会报 1049 Unknown database
        conn = pymysql.connect(
            **{k: v for k, v in config.MYSQL.items() if k != "database"},
            connect_timeout=3,
        )
    except Exception as exc:                    # 包括 2003 连接失败、1045 密码错误
        logger.warning("MySQL 连接失败（%s），已降级为导出 SQL 脚本", str(exc)[:90])
        return False

    sql = """
    INSERT INTO books (upc, title, category, price, price_excl_tax, tax, rating,
                       stock, reviews, description, image_url, detail_url)
    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
    ON DUPLICATE KEY UPDATE
        price=VALUES(price), stock=VALUES(stock), rating=VALUES(rating),
        description=VALUES(description), detail_url=VALUES(detail_url)
    """
    try:
        with conn.cursor() as cur:
            cur.execute(f"CREATE DATABASE IF NOT EXISTS {config.MYSQL['database']} "
                        "DEFAULT CHARSET utf8mb4")
            cur.execute(f"USE {config.MYSQL['database']}")
            cur.execute(CREATE_TABLE_SQL)
            rows = [tuple(r[f] for f in FIELDS) for r in records]
            cur.executemany(sql, rows)
        conn.commit()
        logger.info("MySQL 入库成功：%s 条", len(records))
        return True
    except MySQLError as exc:
        logger.error("MySQL 写入失败：%s", exc)
        return False
    finally:
        conn.close()
